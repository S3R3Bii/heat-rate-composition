"""Independent openpyxl identities and lxml tokens for the complete fixed v2 source."""

import argparse, csv, gzip, hashlib, io, json, posixpath
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from lxml import etree
from openpyxl.utils import get_column_letter, column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
FIELDS = [
    "year",
    "plant",
    "chp",
    "state",
    "prime_mover",
    "fuel",
    "nuclear_unit",
    "elec_mmbtu",
    "net_mwh",
]


def digest(b):
    return hashlib.sha256(b).hexdigest()


def raw_rows(book):
    # Separate parser/library and shared-string text selection from production.
    with ZipFile(io.BytesIO(book)) as z:
        strings = []
        if "xl/sharedStrings.xml" in z.namelist():
            tree = etree.fromstring(z.read("xl/sharedStrings.xml"))
            strings = ["".join(n.xpath(".//s:t/text()", namespaces=NS)) for n in tree]
        sheets = etree.fromstring(z.read("xl/workbook.xml"))
        links = {
            n.get("Id"): n.get("Target")
            for n in etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        }
        target = next(
            n
            for n in sheets.find("s:sheets", namespaces=NS)
            if n.get("name") == "Page 1 Generation and Fuel Data"
        )
        member = links[target.get(REL)]
        member = (
            member.lstrip("/")
            if member.startswith("/")
            else posixpath.normpath("xl/" + member)
        )
        with z.open(member) as handle:
            for event, row in etree.iterparse(
                handle, events=["end"], tag="{" + NS["s"] + "}row"
            ):
                cells = {}
                for cell in row:
                    ref = cell.get("r")
                    kind = cell.get("t", "n")
                    vv = cell.find("s:v", namespaces=NS)
                    token = "" if vv is None else vv.text or ""
                    value = (
                        strings[int(token)]
                        if kind == "s" and token
                        else (
                            "".join(cell.xpath(".//s:is/s:t/text()", namespaces=NS))
                            if kind == "inlineStr"
                            else token
                        )
                    )
                    cells["".join(c for c in ref if c.isalpha())] = (value, kind)
                yield int(row.get("r")), cells
                row.clear()
                while row.getprevious() is not None:
                    del row.getparent()[0]


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--originals", type=Path, required=True)
    p.add_argument("--public", type=Path, required=True)
    a = p.parse_args()
    cfg = json.loads((ROOT / "protocol_v2.json").read_bytes())
    archives = json.loads((ROOT / "data/source/archives.json").read_bytes())
    manifest = json.loads((a.public / "source_manifest.json").read_bytes())
    compressed = (a.public / "source_cells.csv.gz").read_bytes()
    assert digest(compressed) == manifest["source_cells_sha256"]
    stored = list(csv.DictReader(io.StringIO(gzip.decompress(compressed).decode())))
    actual = {(int(r["year"]), int(r["source_row"])): r for r in stored}
    assert len(actual) == len(stored) == manifest["retained_rows"]
    selected = set(
        cfg["coal_codes"] + cfg["natural_gas_codes"] + cfg["diagnostic_fuels"]
    )
    contexts = defaultdict(lambda: defaultdict(list))
    keys = defaultdict(list)
    native = {}
    columns = {}
    expected_locators = set()
    field_checks = numeric_checks = 0
    maximum_error = Decimal(0)
    maximum_fraction = Decimal(0)
    book_hashes = {}
    for arc in archives:
        raw = (a.originals / f"f923_{arc['year']}.zip").read_bytes()
        assert len(raw) == arc["bytes"] and digest(raw) == arc["sha256"]
        with ZipFile(io.BytesIO(raw)) as z:
            book = z.read(arc["selected_member"])
        book_hashes[arc["year"]] = digest(book)
        workbook = openpyxl.load_workbook(
            io.BytesIO(book), read_only=True, data_only=True
        )
        sheet = workbook[cfg["source_sheet"]]
        header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
        mapped = {}
        for field, label in cfg["source_fields"].items():
            matches = [i for i, v in enumerate(header) if v == label]
            assert len(matches) == 1
            mapped[field] = matches[0]
        columns[arc["year"]] = {k: get_column_letter(v + 1) for k, v in mapped.items()}
        meta = next(m for m in manifest["years"] if m["year"] == arc["year"])
        assert (
            columns[arc["year"]] == meta["columns"]
            and book_hashes[arc["year"]] == meta["workbook_sha256"]
        )
        for rownum, row in enumerate(
            sheet.iter_rows(min_row=7, values_only=True), start=7
        ):

            def value(field):
                v = row[mapped[field]] if mapped[field] < len(row) else None
                return "" if v is None else v

            fuel = value("fuel")
            try:
                plant = Decimal(str(value("plant")))
                year = Decimal(str(value("year")))
            except InvalidOperation:
                assert fuel not in selected
                continue
            if (
                not plant.is_finite()
                or not year.is_finite()
                or plant <= 0
                or plant != plant.to_integral_value()
                or year != arc["year"]
            ):
                assert fuel not in selected
                continue
            plant = str(int(plant))
            state, chp = value("state"), value("chp")
            contexts[arc["year"], plant][state, chp].append(rownum)
            if fuel in selected:
                pm = value("prime_mover")
                assert pm in cfg["allowed_prime_movers"]
                keys[arc["year"], plant, pm, fuel].append(rownum)
                native[arc["year"], rownum] = {f: value(f) for f in FIELDS}
                field_checks += 9
        workbook.close()
        print(
            json.dumps(dict(stage="independent_workbook_identities", year=arc["year"])),
            flush=True,
        )
    quarantine = set()
    events = []
    for (year, plant), mapping in sorted(contexts.items()):
        if len(mapping) > 1:
            quarantine.add(plant)
            events.append(
                dict(
                    plant=plant,
                    year=year,
                    reason="conflicting_plant_year_state_or_chp",
                    contexts=[
                        dict(state=s, chp=c, source_rows=rows)
                        for (s, c), rows in sorted(mapping.items())
                    ],
                )
            )
    for (year, plant, pm, fuel), rows in sorted(keys.items()):
        if len(rows) > 1:
            quarantine.add(plant)
            events.append(
                dict(
                    plant=plant,
                    year=year,
                    reason="duplicate_selected_key",
                    prime_mover=pm,
                    fuel=fuel,
                    source_rows=rows,
                )
            )
    assert (
        sorted(quarantine, key=int) == manifest["quarantined_plants"]
        and events == manifest["quarantine_events"]
    )
    assert len(quarantine) == manifest["quarantined_plant_count"]
    for locator, row in native.items():
        if str(int(Decimal(str(row["plant"])))) not in quarantine:
            expected_locators.add(locator)
    assert set(actual) == expected_locators
    retained_counts=Counter((int(r['year']),r['fuel']) for r in stored)
    assert manifest['retained_rows_by_year_fuel']==[dict(year=y,fuel=f,rows=n) for (y,f),n in sorted(retained_counts.items())]
    excluded_counts=Counter(year for (year,rownum),r in native.items() if str(int(Decimal(str(r['plant'])))) in quarantine)
    assert manifest['quarantined_selected_rows_by_year']==[dict(year=y,rows=n) for y,n in sorted(excluded_counts.items())]
    for arc in archives:
        with ZipFile(a.originals / f"f923_{arc['year']}.zip") as z:
            book = z.read(arc["selected_member"])
        for rownum, cells in raw_rows(book):
            locator = arc["year"], rownum
            if locator not in actual:
                continue
            recorded = actual[locator]
            decoded = native[locator]
            cc = columns[arc["year"]]
            raw = [cells.get(cc[f], ("", ""))[0] for f in FIELDS]
            expected = dict(
                year=str(arc["year"]),
                source_row=str(rownum),
                plant=str(int(Decimal(raw[1]))),
                chp=raw[2],
                state=raw[3],
                prime_mover=raw[4],
                fuel=raw[5],
                nuclear_unit=raw[6],
                elec_mmbtu_raw=raw[7],
                net_mwh_raw=raw[8],
                fuel_xml_type=cells.get(cc["elec_mmbtu"], ("", ""))[1],
                generation_xml_type=cells.get(cc["net_mwh"], ("", ""))[1],
                selected_fields_sha256=digest("\x1f".join(raw).encode()),
            )
            assert recorded == expected
            field_checks += len(expected)
            assert Decimal(str(decoded["plant"])) == Decimal(raw[1]) and Decimal(
                str(decoded["year"])
            ) == Decimal(raw[0])
            field_checks += 2
            for field, text in zip(FIELDS[2:7], raw[2:7]):
                d = decoded[field]
                assert (
                    str(d) == text
                    or d == ""
                    and text == ""
                    or isinstance(d, (int, float))
                    and Decimal(str(d)) == Decimal(text)
                )
                field_checks += 1
            for field, text in [("elec_mmbtu", raw[7]), ("net_mwh", raw[8])]:
                d = decoded[field]
                if isinstance(d, (int, float)):
                    with localcontext() as context:
                        context.prec = 80
                        number = Decimal(text)
                        other = Decimal(str(d))
                        assert number.is_finite() and other.is_finite()
                        error = abs(number - other)
                        limit = Decimal("1e-6") + abs(number) * Decimal("1e-15")
                        assert error <= limit
                        maximum_error = max(maximum_error, error)
                        maximum_fraction = max(maximum_fraction, error / limit)
                        numeric_checks += 1
                else:
                    assert str(d) == text or d == "" and text == ""
                    field_checks += 1
        print(
            json.dumps(
                dict(
                    stage="independent_raw_tokens_and_conversion_bounds",
                    year=arc["year"],
                )
            ),
            flush=True,
        )
    result = dict(
        status="PASS_INDEPENDENT_V2_NATIVE_SOURCE",
        retained_rows=len(stored),
        native_selected_rows=len(native),
        quarantined_selected_rows=sum(excluded_counts.values()),
        quarantined_plants=len(quarantine),
        quarantine_events=len(events),
        selected_native_identity_fields=len(native) * 9,
        source_and_projected_field_conditions=field_checks,
        independent_numeric_conversion_conditions=numeric_checks,
        maximum_numeric_conversion_error=str(maximum_error),
        maximum_fraction_of_declared_conversion_bound=str(maximum_fraction),
        public_source_sha256=manifest["source_cells_sha256"],
        public_manifest_sha256=digest((a.public / "source_manifest.json").read_bytes()),
        readers=dict(
            openpyxl=openpyxl.__version__, lxml=__import__("lxml").__version__
        ),
        ambiguous_quantities_aggregated=False,
        original_v1_gate="FAILED_UNCHANGED",
        decomposition_outcomes_computed=False,
        method="Independent complete openpyxl identity/context/key scan and whole-plant quarantine; separate lxml raw-token/source-hash reconciliation of every retained cell; Decimal checks of workbook-reader conversion. No production parser or quarantine helper imports.",
    )
    (a.public / "independent_source_verification.json").write_bytes(
        (json.dumps(result, indent=2) + "\n").encode()
    )
    print(json.dumps(result))


if __name__ == "__main__":
    main()
